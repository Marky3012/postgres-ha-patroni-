from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from lib.models import Campaign

HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Arial", size=10)
BODY_FONT = Font(name="Arial", size=10)
STATUS_FILL = {
    "PASS": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
    "FAIL": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "ERROR": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
    "WARN": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
    "SKIP": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
}


def _header_row(ws, headers, row=1):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def generate(campaign: Campaign, out_path: str):
    summary = campaign.summary()
    wb = Workbook()

    # --- Summary sheet ---
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "PostgreSQL HA Validation Report"
    ws["A1"].font = Font(bold=True, size=14, name="Arial")
    ws["A2"] = f"Scope: {campaign.config_summary.get('scope', '')}"
    ws["A2"].font = BODY_FONT
    ws["A3"] = f"Duration: {summary['duration_s']}s"
    ws["A3"].font = BODY_FONT

    _header_row(ws, ["Metric", "Value"], row=5)
    rows = [
        ("HA Readiness", summary["readiness_label"]),
        ("Readiness Score", f"{summary['readiness_score']}%"),
        ("Total Tests", summary["total_tests"]),
        ("Passed", summary["counts"]["PASS"]),
        ("Failed", summary["counts"]["FAIL"] + summary["counts"]["ERROR"]),
        ("Warnings", summary["counts"]["WARN"]),
        ("Skipped", summary["counts"]["SKIP"]),
    ]
    for i, (k, v) in enumerate(rows, start=6):
        ws.cell(row=i, column=1, value=k).font = BODY_FONT
        ws.cell(row=i, column=2, value=v).font = BODY_FONT
    _autosize(ws, [22, 20])

    # --- Category status sheet ---
    ws2 = wb.create_sheet("Category Status")
    _header_row(ws2, ["Category", "Status"])
    for i, (cat, status) in enumerate(summary["category_status"].items(), start=2):
        ws2.cell(row=i, column=1, value=cat).font = BODY_FONT
        cell = ws2.cell(row=i, column=2, value=status)
        cell.font = BODY_FONT
        cell.fill = STATUS_FILL.get(status, STATUS_FILL["SKIP"])
    _autosize(ws2, [26, 14])

    # --- Detailed results sheet ---
    ws3 = wb.create_sheet("Test Results")
    headers = ["Category", "Test", "Status", "Duration (s)", "SLA Target", "SLA Actual", "SLA Unit", "SLA Met", "Details", "Metrics"]
    _header_row(ws3, headers)
    for i, r in enumerate(campaign.results, start=2):
        vals = [
            r.category, r.name, r.status.value, r.duration_s,
            r.sla_target, r.sla_actual, r.sla_unit, r.sla_met,
            r.details, ", ".join(f"{k}={v}" for k, v in (r.metrics or {}).items()),
        ]
        for col, v in enumerate(vals, start=1):
            cell = ws3.cell(row=i, column=col, value=v)
            cell.font = BODY_FONT
            if col == 3:
                cell.fill = STATUS_FILL.get(r.status.value, STATUS_FILL["SKIP"])
    _autosize(ws3, [18, 42, 10, 12, 10, 10, 8, 9, 48, 40])

    # --- SLA compliance sheet ---
    ws4 = wb.create_sheet("SLA Compliance")
    _header_row(ws4, ["Test", "Category", "Target", "Actual", "Unit", "Met?"])
    row = 2
    for r in campaign.results:
        if r.sla_target is None:
            continue
        ws4.cell(row=row, column=1, value=r.name).font = BODY_FONT
        ws4.cell(row=row, column=2, value=r.category).font = BODY_FONT
        ws4.cell(row=row, column=3, value=r.sla_target).font = BODY_FONT
        ws4.cell(row=row, column=4, value=r.sla_actual).font = BODY_FONT
        ws4.cell(row=row, column=5, value=r.sla_unit).font = BODY_FONT
        met_cell = ws4.cell(row=row, column=6, value="YES" if r.sla_met else "NO")
        met_cell.font = BODY_FONT
        met_cell.fill = STATUS_FILL["PASS"] if r.sla_met else STATUS_FILL["FAIL"]
        row += 1
    _autosize(ws4, [40, 18, 10, 10, 8, 8])

    wb.save(out_path)
    return out_path
